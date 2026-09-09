#!/usr/bin/env python3
"""
One-off export: reads "docs/Class 910_Traits  Weightages (1).xlsx" and writes the
scoring engine's + counsellor chart's reference data as JSON fixtures into
prisma/seed-data/assessment-scoring/. JSON (not TS) because this data now lives in the
database (see src/modules/assessment/scoring/data/store.ts) — these fixtures are only
read once, by prisma/seed-scoring.ts, not compiled into the app's runtime.

Outputs (prisma/seed-data/assessment-scoring/):
  trait-definitions.json          - 18 traits: layer, key, traitName, description,
                                     studentQuality, studentFriendlyExplanation
  riasec-120.json                 - 120 ordered 3-letter RIASEC codes -> style + descriptions
  bigfive-20.json                 - 20 ordered 2-letter Big Five codes -> style + descriptions
  stream-weights.json             - Class 11&12 sub-streams, top-5 trait weights (sum 100)
  domain-weights.json             - career-library industries/domains, top-5 trait weights
  graduate-streams.json           - graduation pathways, top-5 trait weights (sum 100)
  reliability-measures.json       - RVS/ARI/ACI/ORI -> friendly name + what it measures
  scri-band-guidance.json         - SCRI bands 1-4 -> label meaning + student/parent tips
  alignment-rating-guidance.json  - AlignmentRating enum values -> student note

Not part of the app's runtime — rerun manually if the source workbook changes, then
`pnpm db:seed:scoring` to load the fixtures into Postgres.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "Class 910_Traits  Weightages (1).xlsx"
OUT_DIR = ROOT / "prisma" / "seed-data" / "assessment-scoring"
OUT_DIR.mkdir(parents=True, exist_ok=True)

# Workbook trait-column header -> canonical trait key (matches AssessmentQuestion.trait).
# "Decision Confidence" is a vestigial always-empty 19th column and is intentionally
# not mapped.
TRAIT_HEADER_TO_KEY = {
    "Realistic": "REALISTIC",
    "Investigative": "INVESTIGATIVE",
    "Artistic": "ARTISTIC",
    "Social": "SOCIAL",
    "Enterprising": "ENTERPRISING",
    "Conventional": "CONVENTIONAL",
    "Openness": "OPENNESS",
    "Conscientiousness": "CONSCIENTIOUSNESS",
    "Extraversion": "EXTRAVERSION",
    "Agreeableness": "AGREEABLENESS",
    "Emotional Stability": "EMOTIONAL_STABILITY",
    "Numerical Reasoning": "NUMERICAL",
    "Verbal Reasoning": "VERBAL",
    "Logical Reasoning": "LOGICAL",
    "Spatial Reasoning": "SPATIAL",
    "Learning Velocity": "LEARNING_VELOCITY",
    "Uncertainty Tolerance": "UNCERTAINTY_TOLERANCE",
    "Autonomy Preference": "AUTONOMY_PREFERENCE",
}

# RIASEC single-letter code -> canonical key (for the 3-letter DCS codes).
RIASEC_LETTER = {
    "R": "REALISTIC",
    "I": "INVESTIGATIVE",
    "A": "ARTISTIC",
    "S": "SOCIAL",
    "E": "ENTERPRISING",
    "C": "CONVENTIONAL",
}

# SCRI sheet "Rating" label -> AlignmentRating enum value (prisma/schema.prisma).
ALIGNMENT_LABEL_TO_ENUM = {
    "Strongly Aligned": "STRONGLY_ALIGNED",
    "Partially Aligned": "PARTIALLY_ALIGNED",
    "Misaligned": "MISALIGNED",
    "Not Yet Assessed": "NOT_YET_ASSESSED",
}


def s(v):
    if v is None:
        return None
    v = str(v).strip()
    return v if v else None


def rows(ws):
    return list(ws.iter_rows(values_only=True))


def write_json(filename, data):
    (OUT_DIR / filename).write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n")
    print(f"  wrote {filename} ({len(data)} rows)")


def find_header_row(all_rows, must_contain):
    for i, r in enumerate(all_rows):
        cells = {s(c) for c in r}
        if must_contain <= cells:
            return i
    raise SystemExit(f"header row containing {must_contain} not found")


def trait_weights(row, header):
    """Extract {traitKey: weight} for the non-empty weight columns of a row."""
    out = {}
    for idx, head in header.items():
        key = TRAIT_HEADER_TO_KEY.get(head)
        if key is None:
            continue
        val = row[idx] if idx < len(row) else None
        if val is None or str(val).strip() == "":
            continue
        out[key] = float(val)
    return out


def build_header_index(header_row):
    return {i: s(c) for i, c in enumerate(header_row) if s(c)}


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    print(f"sheets: {wb.sheetnames}")

    # ---- Trait Description (traits table + reliability-measures table) ----
    ws = wb["Trait Description"]
    tr = rows(ws)
    trait_defs = []
    measures_start = None
    for i, r in enumerate(tr[1:], start=1):
        layer, trait, name, desc = s(r[0]), s(r[1]), s(r[2]), s(r[3])
        if not layer:
            measures_start = i + 1  # skip the blank separator row
            break
        key = TRAIT_HEADER_TO_KEY.get(trait) or TRAIT_HEADER_TO_KEY.get(f"{trait} Reasoning")
        if key is None:
            # RIASEC/BigFive names match directly; aptitude names lack "Reasoning" suffix
            key = {
                "Numerical Reasoning": "NUMERICAL",
                "Verbal Reasoning": "VERBAL",
                "Logical Reasoning": "LOGICAL",
                "Spatial Reasoning": "SPATIAL",
            }.get(trait)
        if key is None:
            raise SystemExit(f"unmapped trait in Trait Description: {trait!r}")
        trait_defs.append(
            {
                "layer": layer,
                "key": key,
                "trait": trait,
                "traitName": name,
                "description": desc,
                "studentQuality": s(r[4]),
                "studentFriendlyExplanation": s(r[5]),
            }
        )
    assert len(trait_defs) == 18, f"expected 18 traits, got {len(trait_defs)}"
    write_json("trait-definitions.json", trait_defs)

    assert measures_start is not None, "reliability-measures table not found"
    header_row = tr[measures_start]
    assert s(header_row[0]) == "Measure" and s(header_row[3]) == "What It Measures", (
        f"unexpected reliability-measures header: {header_row!r}"
    )
    measures = []
    for r in tr[measures_start + 1 :]:
        measure = s(r[0])
        if not measure:
            continue
        m = re.search(r"\(([A-Z]+)\)\s*$", measure)
        if not m:
            raise SystemExit(f"can't derive a short code from measure {measure!r}")
        measures.append(
            {
                "code": m.group(1),
                "measure": measure,
                "friendlyName": s(r[1]),
                "whatItMeasures": s(r[3]),
            }
        )
    assert len(measures) == 4, f"expected 4 reliability measures, got {len(measures)}"
    write_json("reliability-measures.json", measures)

    # ---- SCRI sheet (band guidance table + academic x career alignment table) ----
    ws = wb["SCRI"]
    scr = rows(ws)
    scri_bands = []
    for i, r in enumerate(scr[1:5], start=1):  # bands 1-4, rows directly under the header
        scri_bands.append(
            {
                "band": i,
                "scoreRange": s(r[0]).split(" (")[0] if s(r[0]) else None,
                "label": s(r[1]),
                "labelMeaning": s(r[2]),
                "forStudents": s(r[3]),
                "tipsForStudents": s(r[4]),
                "tipsForParent": s(r[5]),
            }
        )
    assert len(scri_bands) == 4, f"expected 4 SCRI bands, got {len(scri_bands)}"
    write_json("scri-band-guidance.json", scri_bands)

    assert s(scr[1][8]) == "Rating" and s(scr[1][9]) == "Student Note", (
        f"unexpected alignment-guidance header row: {scr[1][8:10]!r}"
    )
    alignment = []
    for r in scr[2:6]:
        label = s(r[8])
        if not label:
            continue
        enum_value = ALIGNMENT_LABEL_TO_ENUM.get(label)
        if enum_value is None:
            raise SystemExit(f"unmapped alignment rating label: {label!r}")
        alignment.append({"rating": enum_value, "label": label, "studentNote": s(r[9])})
    assert len(alignment) == 4, f"expected 4 alignment ratings, got {len(alignment)}"
    write_json("alignment-rating-guidance.json", alignment)

    # ---- RIASEC 120 ----
    ws = wb["RIASEC 120"]
    rr = rows(ws)
    hi = find_header_row(rr, {"Code (Rank1-Rank2-Rank3)", "Dominant Career Style"})
    header = build_header_index(rr[hi])
    col = {v: k for k, v in header.items()}
    riasec = []
    for r in rr[hi + 1 :]:
        code = s(r[col["Code (Rank1-Rank2-Rank3)"]])
        if not code:
            continue
        riasec.append(
            {
                "code": code,
                "traits": [RIASEC_LETTER[c] for c in code],
                "style": s(r[col["Dominant Career Style"]]),
                "description": s(r[col["Description"]]),
                "explanation": s(r[col["Student & Parent-Friendly Explanation"]]),
            }
        )
    assert len(riasec) == 120, f"expected 120 RIASEC codes, got {len(riasec)}"
    write_json("riasec-120.json", riasec)

    # ---- Big Five 20 ----
    ws = wb["Big Five 20"]
    br = rows(ws)
    hi = find_header_row(br, {"Code (Rank1-Rank2)", "Personality Style"})
    header = build_header_index(br[hi])
    col = {v: k for k, v in header.items()}
    bigfive = []
    for r in br[hi + 1 :]:
        code = s(r[col["Code (Rank1-Rank2)"]])
        if not code:
            continue
        bigfive.append(
            {
                "code": code,  # e.g. "O-C"
                "style": s(r[col["Personality Style"]]),
                "description": s(r[col["Description"]]),
                "explanation": s(r[col["Student & Parent-Friendly Explanation"]]),
            }
        )
    assert len(bigfive) == 20, f"expected 20 Big Five codes, got {len(bigfive)}"
    write_json("bigfive-20.json", bigfive)

    # ---- Class 11&12 Stream weights ----
    ws = wb["Class 11&12_Stream"]
    sr = rows(ws)
    hi = find_header_row(sr, {"Main Stream", "Sub-Streams"})
    header = build_header_index(sr[hi])
    col = {v: k for k, v in header.items()}
    streams = []
    for r in sr[hi + 1 :]:
        sub = s(r[col["Sub-Streams"]])
        if not sub:
            continue
        w = trait_weights(r, header)
        total = round(sum(w.values()), 2)
        assert total == 100, f"stream {sub!r} weights sum to {total}, not 100"
        streams.append(
            {
                "mainStream": s(r[col["Main Stream"]]),
                "subStream": sub,
                "coreSubjects": s(r[col["Core Subjects Usually Offered"]]),
                "electiveSubjects": s(r[col.get("Optional / Elective Subjects", -1)])
                if "Optional / Elective Subjects" in col
                else None,
                "explanation": s(r[col["Student & Parent-Friendly Explanation"]]),
                "weights": w,
            }
        )
    write_json("stream-weights.json", streams)

    # ---- Domain weights (Career Matching) ----
    # Keyed by (industry, domain). Most industries have a single "All Domains" row that
    # covers every domain under them; Defence / Merchant Navy / Entrepreneurship instead
    # enumerate specific domains (and a few of those rows sum to 85-95, not 100 — the
    # engine normalizes by weightSum).
    ws = wb["Domain Wtg"]
    dr = rows(ws)
    hi = find_header_row(dr, {"Career Cluster", "Industry"})
    header = build_header_index(dr[hi])
    col = {v: k for k, v in header.items()}
    domains = []
    for r in dr[hi + 1 :]:
        ind = s(r[col["Industry"]])
        if not ind:
            continue
        w = trait_weights(r, header)
        total = round(sum(w.values()), 2)
        if total != 100:
            print(f"  WARN domain weights {ind!r}/{s(r[col['Domain']])!r} sum to {total}, not 100")
        domains.append(
            {
                "cluster": s(r[col["Career Cluster"]]),
                "industry": ind,
                "domain": s(r[col["Domain"]]),  # "All Domains" or a specific domain
                "weightSum": total,
                "explanation": s(r[col["Student & Parent-Friendly Explanation"]]),
                "weights": w,
            }
        )
    write_json("domain-weights.json", domains)

    # ---- Graduate stream weights (Graduation Pathways) ----
    ws = wb["Graduate_Streams"]
    gr = rows(ws)
    hi = find_header_row(gr, {"Main Stream", "Sub-Streams"})
    header = build_header_index(gr[hi])
    col = {v: k for k, v in header.items()}
    grads = []
    for r in gr[hi + 1 :]:
        sub = s(r[col["Sub-Streams"]])
        if not sub:
            continue
        w = trait_weights(r, header)
        total = round(sum(w.values()), 2)
        assert total == 100, f"graduate stream {sub!r} weights sum to {total}, not 100"
        grads.append(
            {
                "clusterHead": s(r[col.get("Cluster Head", -1)]) if "Cluster Head" in col else None,
                "mainStream": s(r[col["Main Stream"]]),
                "subStream": sub,
                "specialisations": s(r[col.get("Common Specialisations", -1)])
                if "Common Specialisations" in col
                else None,
                "eligibility": s(r[col.get("Eligibility from Class 12", -1)])
                if "Eligibility from Class 12" in col
                else None,
                "keyExams": s(r[col.get("Key Exams To Watch For", -1)])
                if "Key Exams To Watch For" in col
                else None,
                "explanation": s(r[col["Student & Parent-Friendly Explanation"]]),
                "weights": w,
            }
        )
    write_json("graduate-streams.json", grads)


if __name__ == "__main__":
    main()
